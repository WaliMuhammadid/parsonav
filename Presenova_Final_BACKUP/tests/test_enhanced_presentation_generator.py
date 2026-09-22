"""
Automated Test Suite for Enhanced Presentation Generator (Gamma-Style Multi-Stage Engine)
Tests:
1. Outline generation with smart layout assignment
2. Programmatic matplotlib chart generation & image embedding
3. Compiling final PPTX from confirmed outline
4. PPTX structure inspection (verifying shapes, slide count, and embedded pictures)
"""

import os
import sys
from pptx import Presentation

# Ensure project root is in path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

from services.presentation_generator import (
    generate_presentation_outline,
    build_pptx_from_outline,
    classify_slide_intent,
    THEMES
)

def test_classify_slide_intent():
    """Verify deterministic intent classification across layout archetypes."""
    # Chart intent
    assert classify_slide_intent("System Performance Trends", ["Observed 85% growth rate"], [], "chart") == "chart"
    # Stat callout intent
    assert classify_slide_intent("Impact Metrics", ["99.9% uptime SLA"], [], "stat") == "stat_callout"
    # Timeline intent
    assert classify_slide_intent("Roadmap", ["Phase 1: Ingestion", "Phase 2: Execution"], [], "timeline") == "timeline"
    # Comparison grid intent
    assert classify_slide_intent("Comparison", [], [{"title": "Legacy"}, {"title": "Presenova"}], "cards") == "comparison_grid"
    # Title intent
    assert classify_slide_intent("Mastering AI", [], [], "title") == "title"


def test_outline_generation():
    """Test generating a structured outline with 5 slides."""
    outline = generate_presentation_outline(
        topic="Autonomous Vehicles & Edge Computing",
        slide_count=5,
        tone="Professional",
        audience="Tech Executives"
    )

    assert outline is not None
    assert "slides" in outline
    assert len(outline["slides"]) == 5
    assert outline["topic"] == "Autonomous Vehicles & Edge Computing"

    # Slide 1 must be title
    assert outline["slides"][0]["type"] == "title"
    # Check that at least one slide has chart_data
    has_chart = any("chart_data" in s or s.get("type") == "chart" for s in outline["slides"])
    assert has_chart, "Expected at least one slide to have chart_data for visual rendering"


def test_build_pptx_with_visual_charts():
    """Test building final PPTX and verifying shapes and embedded chart pictures."""
    outline = generate_presentation_outline(
        topic="Quantum Computing Horizons",
        slide_count=5,
        tone="Academic",
        audience="Physics Faculty"
    )

    pptx_path = build_pptx_from_outline(
        outline_data=outline,
        theme_name="academic_elegant",
        custom_overrides={
            "show_slide_numbers": True,
            "confidentiality_tag": "Presenova Research 2026"
        }
    )

    assert os.path.exists(pptx_path), f"File {pptx_path} does not exist"

    # Inspect generated presentation
    prs = Presentation(pptx_path)
    assert len(prs.slides) == 5

    # Check for embedded picture (the visual chart generated via matplotlib)
    picture_count = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE
                picture_count += 1

    assert picture_count >= 1, f"Expected at least 1 embedded visual chart picture, found {picture_count}"

    # Verify slide notes
    has_notes = any(slide.has_notes_slide and slide.notes_slide.notes_text_frame.text for slide in prs.slides)
    assert has_notes, "Expected speaker notes to be populated"

    print(f"\n[PASS] PPTX generated successfully with {picture_count} embedded visual chart(s) at: {pptx_path}")

if __name__ == '__main__':
    test_classify_slide_intent()
    test_outline_generation()
    test_build_pptx_with_visual_charts()
    print("\nAll 3 test suites passed successfully!")
